import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
OUTPUT_PATH = DATA_DIR / "openapi.json"


def parse_type(type_value: str) -> tuple[str, bool]:
    normalized = (type_value or "").upper()
    required = "필수" in normalized
    schema_type = "integer" if "INTEGER" in normalized else "string"
    return schema_type, required


def make_parameter(name: str, type_value: str, title: str, description: str) -> dict:
    schema_type, required = parse_type(type_value)
    parameter = {
        "name": name,
        "in": "query",
        "required": required,
        "description": description or title,
        "schema": {
            "type": schema_type,
        },
    }

    if name == "Key":
        parameter["schema"]["default"] = "sample"
    elif name == "Type":
        parameter["schema"]["default"] = "xml"
        parameter["schema"]["enum"] = ["xml", "json"]
    elif name == "pIndex":
        parameter["schema"]["default"] = 1
        parameter["schema"]["minimum"] = 1
    elif name == "pSize":
        parameter["schema"]["default"] = 100
        parameter["schema"]["minimum"] = 1
        parameter["schema"]["maximum"] = 1000

    return parameter


def parse_sheet(workbook_path: Path) -> dict:
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    title = None
    request_url = None
    current_section = None
    basic_params = []
    request_params = []
    output_fields = []
    messages = []

    for row in sheet.iter_rows(values_only=True):
        cells = list(row)
        first = cells[0]

        if isinstance(first, str) and first.startswith("(") and first.endswith(")"):
            title = first.strip("()")
            continue

        if first == "• 요청주소":
            current_section = "url"
            continue
        if isinstance(first, str) and first.startswith("- https://open.neis.go.kr/hub/"):
            request_url = first[2:]
            continue
        if first == "• 기본인자":
            current_section = "basic"
            continue
        if first == "• 요청인자":
            current_section = "request"
            continue
        if first == "• 출력값":
            current_section = "output"
            continue
        if first == "• 메시지 설명":
            current_section = "message"
            continue

        if first is None:
            continue

        if current_section in {"basic", "request"} and first != "변수명":
            target = basic_params if current_section == "basic" else request_params
            target.append(
                {
                    "name": cells[0],
                    "type": cells[1],
                    "title": cells[2],
                    "description": cells[3],
                }
            )
        elif current_section == "output" and first != "No":
            output_fields.append(
                {
                    "name": cells[1],
                    "title": cells[2],
                    "description": cells[3],
                }
            )
        elif current_section == "message" and first != "No":
            messages.append(
                {
                    "category": cells[1],
                    "code": str(cells[2]),
                    "description": cells[3],
                }
            )

    return {
        "title": title,
        "url": request_url,
        "basic_params": basic_params,
        "request_params": request_params,
        "output_fields": output_fields,
        "messages": messages,
    }


def build_operation_id(path_name: str) -> str:
    return f"get{path_name[:1].upper()}{path_name[1:]}"


def build_spec(api_sheets: list[dict]) -> dict:
    common_messages = api_sheets[0]["messages"] if api_sheets else []
    message_description = "\n".join(
        f"- {message['category']} {message['code']}: {message['description']}"
        for message in common_messages
    )

    spec = {
        "openapi": "3.0.4",
        "info": {
            "title": "NEIS Open API",
            "version": "1.0.0",
            "description": "NEIS 공개 API 엑셀 명세를 기반으로 생성한 OpenAPI 문서입니다.",
        },
        "servers": [
            {
                "url": "https://open.neis.go.kr",
                "description": "NEIS Open API",
            }
        ],
        "paths": {},
        "components": {
            "schemas": {},
        },
    }

    for api in api_sheets:
        url_path = "/" + api["url"].split("https://open.neis.go.kr", 1)[1].lstrip("/")
        path_name = url_path.rsplit("/", 1)[-1]
        record_schema_name = f"{path_name[:1].upper()}{path_name[1:]}Record"
        response_schema_name = f"{path_name[:1].upper()}{path_name[1:]}Response"

        parameters = [
            make_parameter(param["name"], param["type"], param["title"], param["description"])
            for param in [*api["basic_params"], *api["request_params"]]
        ]

        record_properties = {}
        for field in api["output_fields"]:
            record_properties[field["name"]] = {
                "type": "string",
                "description": field["description"] or field["title"],
            }

        spec["components"]["schemas"][record_schema_name] = {
            "type": "object",
            "description": f"{api['title']} 단건 데이터",
            "properties": record_properties,
        }

        spec["components"]["schemas"][response_schema_name] = {
            "type": "object",
            "description": f"{api['title']} 응답",
            "properties": {
                path_name: {
                    "type": "array",
                    "description": f"{api['title']} 데이터 목록",
                    "items": {
                        "$ref": f"#/components/schemas/{record_schema_name}"
                    },
                }
            },
        }

        spec["paths"][url_path] = {
            "get": {
                "summary": f"{api['title']} 조회",
                "description": message_description,
                "operationId": build_operation_id(path_name),
                "parameters": parameters,
                "responses": {
                    "200": {
                        "description": "정상 응답",
                        "content": {
                            "application/json": {
                                "schema": {
                                    "$ref": f"#/components/schemas/{response_schema_name}"
                                }
                            },
                            "application/xml": {
                                "schema": {
                                    "type": "string"
                                }
                            },
                        },
                    },
                    "400": {
                        "description": "잘못된 요청 또는 필수값 누락",
                    },
                    "401": {
                        "description": "유효하지 않은 인증키",
                    },
                    "500": {
                        "description": "서버 또는 데이터베이스 오류",
                    },
                },
                "tags": ["NEIS"],
            }
        }

    return spec


def main() -> None:
    workbooks = sorted(DATA_DIR.glob("*오픈API명세서.xlsx"))
    api_sheets = [parse_sheet(path) for path in workbooks]
    spec = build_spec(api_sheets)
    OUTPUT_PATH.write_text(
        json.dumps(spec, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
